import io

import pytest
from docx import Document
from fastapi.testclient import TestClient

import main
from db import get_store
from db.seed import seed_if_empty
from db.store import Store
from models.journal import HardConstraint, JournalRequirementSpec
from routers.match import get_embedding_provider
from services.journal_lists.service import MAX_CSV_BYTES, normalize_issn, parse_csv


class FakeEmbeddings:
    VOCAB = ["medical", "imaging", "x-ray", "deep", "learning", "agriculture", "crop"]

    def encode(self, texts):
        return [[float(word in t.lower()) + 0.01 for word in self.VOCAB] for t in texts]


@pytest.fixture
def store(tmp_path):
    s = Store(tmp_path / "lists.db")
    seed_if_empty(s)
    return s


@pytest.fixture
def client(store):
    main.app.dependency_overrides[get_store] = lambda: store
    main.app.dependency_overrides[get_embedding_provider] = lambda: FakeEmbeddings()
    yield TestClient(main.app)
    main.app.dependency_overrides.clear()


def docx_bytes() -> bytes:
    doc = Document()
    doc.add_paragraph("Deep Learning for Medical Imaging of X-ray Fractures")
    doc.add_paragraph("Abstract")
    doc.add_paragraph("We apply deep learning to medical imaging of x-ray data.")
    doc.add_paragraph("Keywords: deep learning, medical imaging")
    doc.add_paragraph("Introduction")
    doc.add_paragraph("Body text about medical imaging.")
    doc.add_paragraph("References")
    doc.add_paragraph("[1] A. Author, A paper, 2024.")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def upload_manuscript(client):
    return client.post(
        "/manuscripts/upload",
        files={"file": ("paper.docx", docx_bytes(), "application/octet-stream")},
    ).json()["manuscript_id"]


def add_fourth_demo_journal(store):
    store.save_journal(
        JournalRequirementSpec(
            journal_id="demo-open-001",
            name="Demo Journal of Open Methods",
            short_name="Demo JOM",
            publisher="Demo Publisher",
            source_url="https://example.com/demo-jom/author-guidelines",
            hard_constraints=HardConstraint(),
            aims="Research on open methods for machine learning and data science.",
            scope_description="Open methods, artificial intelligence, medical imaging, and reproducible research.",
            topics=["artificial intelligence", "medical imaging", "open methods"],
            accepted_article_types=["research_article", "review"],
            languages=["en"],
            access_model="open_access",
            apc_usd=500,
            review_speed_days_avg=30,
            indexes=["scopus"],
            extraction_confidence=1,
            needs_human_review=False,
            is_demo=True,
        )
    )


def create_four_resolved_list(client):
    return client.post(
        "/journal-lists",
        json={
            "name": "Four resolved journals",
            "journals": [
                {"journal_id": "demo-ai-001"},
                {"journal_id": "demo-med-001"},
                {"journal_id": "demo-agri-001"},
                {"journal_id": "demo-open-001"},
            ],
        },
    ).json()["journal_list"]["journal_list_id"]


def test_parse_csv_extra_columns_unicode_and_issn():
    rows, warnings = parse_csv(
        "Journal Name,ISSN,Reward Eligible,Reward Amount\n"
        "Demo Journal of Medical Imaging,1234567X,Yes,5000\n"
        "مجلة عربية,,No,\n".encode("utf-8")
    )

    assert warnings == []
    assert rows[0]["journal_name"] == "Demo Journal of Medical Imaging"
    assert rows[0]["metadata"]["reward eligible"] == "Yes"
    assert normalize_issn(rows[0]["issn"]) == "1234-567X"
    assert rows[1]["journal_name"] == "مجلة عربية"


def test_create_retrieve_and_delete_journal_list(client):
    response = client.post(
        "/journal-lists",
        json={
            "name": "IAU Approved Journals 2026",
            "institution": "IAU",
            "journals": [
                {"journal_name": "Demo Journal of Medical Imaging", "reward_eligible": "Yes"},
                {"journal_name": "Unknown Journal", "notes": "Needs acquisition"},
            ],
        },
    )

    assert response.status_code == 200
    body = response.json()
    jid = body["journal_list"]["journal_list_id"]
    assert body["journal_list"]["journal_count"] == 2
    assert body["journal_list"]["resolved_count"] == 1
    assert body["journal_list"]["unresolved_count"] == 1
    assert body["entries"][0]["resolution_status"] == "resolved"
    assert body["entries"][0]["metadata"]["reward_eligible"] == "Yes"
    assert body["entries"][1]["resolution_status"] == "unresolved"

    assert client.get(f"/journal-lists/{jid}").json()["journal_list"]["name"] == "IAU Approved Journals 2026"
    assert len(client.get("/journal-lists").json()) == 1
    assert client.delete(f"/journal-lists/{jid}").status_code == 204
    assert client.get(f"/journal-lists/{jid}").status_code == 404


def test_upload_csv_marks_duplicates_invalid_and_unknown(client):
    content = (
        "Journal Name,ISSN,Reward Eligible\n"
        "Demo Journal of Medical Imaging,,Yes\n"
        "Demo Journal of Medical Imaging,,Yes\n"
        ",,\n"
        "Not In Warraq,,No\n"
    ).encode()
    response = client.post(
        "/journal-lists/upload",
        data={"name": "Reward list"},
        files={"file": ("journals.csv", content, "text/csv")},
    )

    assert response.status_code == 200
    statuses = [entry["resolution_status"] for entry in response.json()["entries"]]
    assert statuses == ["resolved", "duplicate", "unresolved"]


def test_upload_csv_rejects_empty_and_no_valid_rows(client):
    empty = client.post(
        "/journal-lists/upload",
        files={"file": ("journals.csv", b"", "text/csv")},
    )
    assert empty.status_code == 400
    assert "empty" in empty.json()["detail"]

    only_blank_rows = client.post(
        "/journal-lists/upload",
        files={"file": ("journals.csv", b"Journal Name,ISSN\n,\n", "text/csv")},
    )
    assert only_blank_rows.status_code == 400
    assert "did not contain any journal rows" in only_blank_rows.json()["detail"]


def test_upload_csv_rejects_unusable_binary_content(client):
    response = client.post(
        "/journal-lists/upload",
        files={"file": ("journals.csv", b"\xff\xfe\x00\x00", "text/csv")},
    )

    assert response.status_code == 400
    assert "binary data" in response.json()["detail"]


def test_upload_csv_rejects_wrong_extension_and_large_file(client):
    wrong_extension = client.post(
        "/journal-lists/upload",
        files={"file": ("journals.xlsx", b"Journal Name\nDemo Journal of Medical Imaging\n", "text/csv")},
    )
    assert wrong_extension.status_code == 400
    assert "CSV" in wrong_extension.json()["detail"]

    too_large = b"Journal Name\n" + b"A" * MAX_CSV_BYTES
    response = client.post(
        "/journal-lists/upload",
        files={"file": ("journals.csv", too_large, "text/csv")},
    )
    assert response.status_code == 400
    assert "larger than 2 MB" in response.json()["detail"]


def test_upload_csv_enforces_500_row_limit(client):
    content = "Journal Name\n" + "\n".join(f"Journal {i}" for i in range(501))
    response = client.post(
        "/journal-lists/upload",
        files={"file": ("journals.csv", content.encode(), "text/csv")},
    )

    assert response.status_code == 400
    assert "limited to 500" in response.json()["detail"]


def test_custom_list_matching_only_returns_provided_resolved_journals(client):
    manuscript_id = upload_manuscript(client)
    created = client.post(
        "/journal-lists",
        json={
            "name": "Two journal institutional list",
            "journals": [
                {"journal_name": "Demo Journal of Medical Imaging", "reward_eligible": "Yes", "category": "A"},
                {"journal_name": "Unknown Journal", "reward_eligible": "No"},
            ],
        },
    ).json()
    list_id = created["journal_list"]["journal_list_id"]

    response = client.post(
        f"/journal-lists/{list_id}/match",
        json={"manuscript_id": manuscript_id, "preferences": {"top_k": 10}},
    )

    assert response.status_code == 200
    body = response.json()
    assert body["provided_count"] == 2
    assert body["resolved_count"] == 1
    assert body["unresolved_count"] == 1
    assert [item["match"]["journal_id"] for item in body["matches"]] == ["demo-med-001"]
    assert body["matches"][0]["entry"]["metadata"]["category"] == "A"
    assert body["unresolved_entries"][0]["original_name"] == "Unknown Journal"
    assert body["unresolved_entries"][0]["resolution_message"]


def test_custom_list_can_resolve_known_journals_by_journal_id(client):
    response = client.post(
        "/journal-lists",
        json={
            "name": "Selected in Warraq",
            "journals": [
                {"journal_id": "demo-ai-001", "journal_name": "Optional display name"},
                {"journal_id": "missing-journal", "journal_name": "Missing"},
            ],
        },
    )

    assert response.status_code == 200
    body = response.json()
    assert body["journal_list"]["resolved_count"] == 1
    assert body["journal_list"]["unresolved_count"] == 1
    assert body["entries"][0]["resolution_status"] == "resolved"
    assert body["entries"][0]["matched_journal_id"] == "demo-ai-001"
    assert "journal_id" in body["entries"][0]["resolution_message"]
    assert body["entries"][1]["resolution_status"] == "unresolved"


def test_custom_list_excludes_highest_similarity_outside_list(client):
    manuscript_id = upload_manuscript(client)
    list_id = client.post(
        "/journal-lists",
        json={
            "name": "Agriculture only",
            "journals": [{"journal_id": "demo-agri-001"}],
        },
    ).json()["journal_list"]["journal_list_id"]

    response = client.post(
        f"/journal-lists/{list_id}/match",
        json={"manuscript_id": manuscript_id, "preferences": {"top_k": 10}},
    )

    assert response.status_code == 200
    body = response.json()
    assert body["eligible_count"] == 1
    assert [item["match"]["journal_id"] for item in body["matches"]] == ["demo-agri-001"]
    assert "demo-med-001" not in {item["match"]["journal_id"] for item in body["matches"]}


def test_custom_list_match_404s_for_missing_list_or_manuscript(client):
    missing_list = client.post(
        "/journal-lists/not-a-list/match",
        json={"manuscript_id": "not-a-manuscript", "preferences": {}},
    )
    assert missing_list.status_code == 404
    assert missing_list.json()["detail"] == "Journal list not found."

    list_id = client.post(
        "/journal-lists",
        json={"name": "Known", "journals": [{"journal_name": "Demo Journal of Medical Imaging"}]},
    ).json()["journal_list"]["journal_list_id"]

    missing_manuscript = client.post(
        f"/journal-lists/{list_id}/match",
        json={"manuscript_id": "not-a-manuscript", "preferences": {}},
    )
    assert missing_manuscript.status_code == 404
    assert missing_manuscript.json()["detail"] == "Manuscript not found. Upload it first."


def test_deleted_journal_list_cannot_be_matched(client):
    manuscript_id = upload_manuscript(client)
    list_id = client.post(
        "/journal-lists",
        json={"name": "Delete me", "journals": [{"journal_name": "Demo Journal of Medical Imaging"}]},
    ).json()["journal_list"]["journal_list_id"]

    assert client.delete(f"/journal-lists/{list_id}").status_code == 204
    response = client.post(
        f"/journal-lists/{list_id}/match",
        json={"manuscript_id": manuscript_id, "preferences": {}},
    )

    assert response.status_code == 404
    assert response.json()["detail"] == "Journal list not found."


def test_custom_list_match_respects_top_k(client):
    manuscript_id = upload_manuscript(client)
    list_id = client.post(
        "/journal-lists",
        json={
            "name": "Two known journals",
            "journals": [
                {"journal_name": "Demo Journal of Medical Imaging"},
                {"journal_name": "Demo Journal of Agricultural Systems"},
            ],
        },
    ).json()["journal_list"]["journal_list_id"]

    response = client.post(
        f"/journal-lists/{list_id}/match",
        json={"manuscript_id": manuscript_id, "preferences": {"top_k": 1}},
    )

    assert response.status_code == 200
    body = response.json()
    assert body["eligible_count"] == 2
    assert len(body["matches"]) == 1


def test_issn_only_entry_is_unresolved_in_v1(client):
    response = client.post(
        "/journal-lists",
        json={"name": "ISSN only", "journals": [{"issn": "1234-567X"}]},
    )

    assert response.status_code == 200
    entry = response.json()["entries"][0]
    assert entry["issn"] == "1234-567X"
    assert entry["resolution_status"] == "unresolved"
    assert "No existing Warraq journal matched" in entry["resolution_message"]


def test_ambiguous_name_resolution_is_reported(store, client):
    duplicate = JournalRequirementSpec(
        journal_id="demo-med-duplicate",
        name="Demo Journal of Medical Imaging",
        publisher="Duplicate Publisher",
        source_url="https://example.test/duplicate",
        hard_constraints=HardConstraint(),
        scope_description="Medical imaging duplicate.",
        extraction_confidence=1,
        topics=["medical imaging"],
        access_model="open_access",
        is_demo=True,
    )
    store.save_journal(duplicate)

    response = client.post(
        "/journal-lists",
        json={"name": "Ambiguous", "journals": [{"journal_name": "Demo Journal of Medical Imaging"}]},
    )

    assert response.status_code == 200
    entry = response.json()["entries"][0]
    assert entry["resolution_status"] == "ambiguous"
    assert entry["matched_journal_id"] is None
    assert "More than one" in entry["resolution_message"]


def test_custom_list_preferences_still_apply(client):
    manuscript_id = upload_manuscript(client)
    list_id = client.post(
        "/journal-lists",
        json={
            "name": "Medical and agriculture only",
            "journals": [
                {"journal_name": "Demo Journal of Medical Imaging"},
                {"journal_name": "Demo Journal of Agricultural Systems"},
            ],
        },
    ).json()["journal_list"]["journal_list_id"]

    response = client.post(
        f"/journal-lists/{list_id}/match",
        json={"manuscript_id": manuscript_id, "preferences": {"required_indexes": ["wos"]}},
    )

    assert response.status_code == 200
    assert {item["match"]["journal_id"] for item in response.json()["matches"]} == {
        "demo-med-001",
        "demo-agri-001",
    }
    assert client.post("/match", json={"manuscript_id": manuscript_id}).status_code == 200


def test_custom_list_with_four_resolved_and_no_restrictive_preferences_matches_all(store, client):
    add_fourth_demo_journal(store)
    manuscript_id = upload_manuscript(client)
    list_id = create_four_resolved_list(client)

    response = client.post(
        f"/journal-lists/{list_id}/match",
        json={"manuscript_id": manuscript_id, "preferences": {"top_k": 10}},
    )

    assert response.status_code == 200
    body = response.json()
    assert body["provided_count"] == 4
    assert body["resolved_count"] == 4
    assert body["unresolved_count"] == 0
    assert body["eligible_count"] == 4
    assert body["excluded_entries"] == []
    assert {item["match"]["journal_id"] for item in body["matches"]} == {
        "demo-ai-001",
        "demo-med-001",
        "demo-agri-001",
        "demo-open-001",
    }


def test_custom_list_preferences_exclude_two_and_rank_remaining_two(store, client):
    add_fourth_demo_journal(store)
    manuscript_id = upload_manuscript(client)
    list_id = create_four_resolved_list(client)

    response = client.post(
        f"/journal-lists/{list_id}/match",
        json={
            "manuscript_id": manuscript_id,
            "preferences": {
                "top_k": 10,
                "open_access_only": True,
                "max_apc": 1000,
            },
        },
    )

    assert response.status_code == 200
    body = response.json()
    assert body["provided_count"] == 4
    assert body["resolved_count"] == 4
    assert body["eligible_count"] == 2
    assert {item["match"]["journal_id"] for item in body["matches"]} == {
        "demo-ai-001",
        "demo-open-001",
    }
    excluded = {item["matched_journal_id"]: item["reasons"] for item in body["excluded_entries"]}
    assert excluded == {
        "demo-med-001": ["max_apc"],
        "demo-agri-001": ["open_access"],
    }


def test_custom_list_preferences_exclude_all_with_clear_diagnostics(store, client):
    add_fourth_demo_journal(store)
    manuscript_id = upload_manuscript(client)
    list_id = create_four_resolved_list(client)

    response = client.post(
        f"/journal-lists/{list_id}/match",
        json={"manuscript_id": manuscript_id, "preferences": {"top_k": 10, "max_review_days": 1}},
    )

    assert response.status_code == 200
    body = response.json()
    assert body["provided_count"] == 4
    assert body["resolved_count"] == 4
    assert body["unresolved_count"] == 0
    assert body["eligible_count"] == 0
    assert body["matches"] == []
    assert {item["matched_journal_id"] for item in body["excluded_entries"]} == {
        "demo-ai-001",
        "demo-med-001",
        "demo-agri-001",
        "demo-open-001",
    }
    assert all(item["reasons"] == ["max_review_days"] for item in body["excluded_entries"])


def test_csv_with_unresolved_entries_only_matches_resolved_journal(client):
    manuscript_id = upload_manuscript(client)
    content = "Journal Name,Notes\nDemo Journal of Medical Imaging,\n" + "\n".join(
        f"Unknown Journal {i}," for i in range(1, 12)
    )
    uploaded = client.post(
        "/journal-lists/upload",
        data={"name": "Mostly unresolved"},
        files={"file": ("journals.csv", content.encode(), "text/csv")},
    )
    assert uploaded.status_code == 200
    list_id = uploaded.json()["journal_list"]["journal_list_id"]

    response = client.post(
        f"/journal-lists/{list_id}/match",
        json={"manuscript_id": manuscript_id, "preferences": {"top_k": 10}},
    )

    assert response.status_code == 200
    body = response.json()
    assert body["provided_count"] == 12
    assert body["resolved_count"] == 1
    assert body["unresolved_count"] == 11
    assert body["eligible_count"] == 1
    assert [item["match"]["journal_id"] for item in body["matches"]] == ["demo-med-001"]
    assert len(body["unresolved_entries"]) == 11
    assert "Unknown Journal 1" in {entry["original_name"] for entry in body["unresolved_entries"]}


def test_unresolved_journals_never_appear_as_match_results(client):
    manuscript_id = upload_manuscript(client)
    created = client.post(
        "/journal-lists",
        json={
            "name": "Resolved and unresolved",
            "journals": [
                {"journal_name": "Demo Journal of Medical Imaging"},
                {"journal_name": "Unknown Journal Outside Warraq"},
            ],
        },
    ).json()
    list_id = created["journal_list"]["journal_list_id"]
    unresolved_name = next(
        entry["original_name"]
        for entry in created["entries"]
        if entry["resolution_status"] == "unresolved"
    )

    response = client.post(
        f"/journal-lists/{list_id}/match",
        json={"manuscript_id": manuscript_id, "preferences": {"top_k": 10}},
    )

    assert response.status_code == 200
    result_names = {item["match"]["name"] for item in response.json()["matches"]}
    assert unresolved_name not in result_names


def test_journal_list_resolution_refreshes_after_journal_is_approved(store, client):
    created = client.post(
        "/journal-lists",
        json={
            "name": "Before approval",
            "journals": [{"journal_name": "Future Approved Journal"}],
        },
    ).json()
    list_id = created["journal_list"]["journal_list_id"]
    assert created["journal_list"]["resolved_count"] == 0
    assert created["journal_list"]["unresolved_count"] == 1
    assert created["entries"][0]["resolution_status"] == "unresolved"

    index_before = client.get("/journal-lists").json()
    indexed_before = next(item for item in index_before if item["journal_list_id"] == list_id)
    assert indexed_before["resolved_count"] == 0
    assert indexed_before["unresolved_count"] == 1

    store.save_journal(
        JournalRequirementSpec(
            journal_id="s1-future-approved",
            name="Future Approved Journal",
            publisher="Future Publisher",
            source_url="https://example.com/future",
            hard_constraints=HardConstraint(),
            aims="Future approved research.",
            scope_description="Medical imaging and artificial intelligence.",
            topics=["medical imaging", "artificial intelligence"],
            accepted_article_types=["research_article"],
            languages=["en"],
            access_model="open_access",
            apc_usd=None,
            review_speed_days_avg=None,
            indexes=[],
            extraction_confidence=1,
            needs_human_review=False,
            is_demo=False,
        )
    )

    refreshed = client.get(f"/journal-lists/{list_id}").json()
    assert refreshed["journal_list"]["resolved_count"] == 1
    assert refreshed["journal_list"]["unresolved_count"] == 0
    assert refreshed["entries"][0]["resolution_status"] == "resolved"
    assert refreshed["entries"][0]["matched_journal_id"] == "s1-future-approved"


def test_journal_list_index_refreshes_after_journal_is_approved(store, client):
    created = client.post(
        "/journal-lists",
        json={
            "name": "Index before approval",
            "journals": [{"journal_name": "Future Indexed Journal"}],
        },
    ).json()
    list_id = created["journal_list"]["journal_list_id"]
    assert created["journal_list"]["resolved_count"] == 0
    assert created["journal_list"]["unresolved_count"] == 1

    store.save_journal(
        JournalRequirementSpec(
            journal_id="s1-future-indexed",
            name="Future Indexed Journal",
            publisher="Future Publisher",
            source_url="https://example.com/future-indexed",
            hard_constraints=HardConstraint(),
            aims="Future indexed research.",
            scope_description="Artificial intelligence and medical imaging.",
            topics=["artificial intelligence", "medical imaging"],
            accepted_article_types=["research_article"],
            languages=["en"],
            access_model="open_access",
            apc_usd=None,
            review_speed_days_avg=None,
            indexes=[],
            extraction_confidence=1,
            needs_human_review=False,
            is_demo=False,
        )
    )

    summaries = client.get("/journal-lists").json()
    refreshed = next(item for item in summaries if item["journal_list_id"] == list_id)
    assert refreshed["resolved_count"] == 1
    assert refreshed["unresolved_count"] == 0

    detail = client.get(f"/journal-lists/{list_id}").json()
    assert detail["journal_list"]["resolved_count"] == refreshed["resolved_count"]
    assert detail["journal_list"]["unresolved_count"] == refreshed["unresolved_count"]
    assert len(detail["entries"]) == 1
    assert detail["entries"][0]["matched_journal_id"] == "s1-future-indexed"
