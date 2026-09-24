"""End-to-end API tests: upload -> journals -> match, against a temporary database."""

import io

import pytest
from docx import Document
from fastapi.testclient import TestClient

import main
from db import get_store
from db.seed import seed_if_empty
from db.store import Store
from models.journal import HardConstraint, JournalRequirementSpec, ReviewQueueItem
from routers.match import get_embedding_provider


class FakeEmbeddings:
    """Keyword-based vectors, so tests don't download a model."""

    VOCAB = ["medical", "imaging", "x-ray", "deep", "learning", "agriculture", "crop"]

    def encode(self, texts):
        return [[float(word in t.lower()) + 0.01 for word in self.VOCAB] for t in texts]


@pytest.fixture
def store(tmp_path):
    s = Store(tmp_path / "test.db")
    seed_if_empty(s)
    return s


@pytest.fixture
def client(store):
    main.app.dependency_overrides[get_store] = lambda: store
    main.app.dependency_overrides[get_embedding_provider] = lambda: FakeEmbeddings()
    yield TestClient(main.app)
    main.app.dependency_overrides.clear()


def docx_bytes(title="Deep Learning for Medical Imaging of X-ray Fractures") -> bytes:
    doc = Document()
    doc.add_paragraph(title)
    doc.add_paragraph("Abstract")
    doc.add_paragraph("We apply deep learning to medical imaging of x-ray data.")
    doc.add_paragraph("Keywords: deep learning, medical imaging")
    doc.add_paragraph("Introduction")
    doc.add_paragraph("Body text about medical imaging.")
    doc.add_paragraph("References")
    doc.add_paragraph("[1] A. Author, A paper, 2024.")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def upload(client, content: bytes, name="paper.docx"):
    return client.post("/manuscripts/upload", files={"file": (name, content, "application/octet-stream")})


def test_upload_saves_and_caches(client):
    first = upload(client, docx_bytes())
    assert first.status_code == 200
    body = first.json()
    assert body["manuscript_id"] and body["from_cache"] is False
    assert body["parsed"]["reference_count"] == 1

    again = upload(client, docx_bytes())
    assert again.json()["from_cache"] is True
    assert again.json()["manuscript_id"] == body["manuscript_id"]

    record = client.get(f"/manuscripts/{body['manuscript_id']}")
    assert record.status_code == 200
    assert record.json()["content_hash"]


def test_upload_rejects_non_docx_and_empty(client):
    assert upload(client, b"%PDF", name="paper.pdf").status_code == 400
    assert upload(client, b"", name="paper.docx").status_code == 400


def test_unknown_ids_return_404(client):
    assert client.get("/manuscripts/nope").status_code == 404
    assert client.get("/journals/nope").status_code == 404
    assert client.post("/match", json={"manuscript_id": "nope"}).status_code == 404


def test_list_and_get_journals(client):
    journals = client.get("/journals").json()
    assert len(journals) == 3
    assert all(j["is_demo"] for j in journals)
    assert journals[0]["requirements_summary"]

    spec = client.get("/journals/demo-med-001").json()
    assert spec["hard_constraints"]["max_abstract_words"] == 300


def test_review_queue(client, store):
    assert client.get("/journals/review-queue").json() == []
    store.save_review_item(
        ReviewQueueItem(journal_id="x", name="X", source_url="https://x", low_confidence_fields=["apc_usd"])
    )
    items = client.get("/journals/review-queue").json()
    assert len(items) == 1 and items[0]["low_confidence_fields"] == ["apc_usd"]


def test_match_by_manuscript_id(client):
    manuscript_id = upload(client, docx_bytes()).json()["manuscript_id"]
    response = client.post("/match", json={"manuscript_id": manuscript_id})
    assert response.status_code == 200
    results = response.json()
    assert [r["rank"] for r in results] == list(range(1, len(results) + 1))
    assert results[0]["journal_id"] == "demo-med-001"
    assert results[0]["scope_fit"] in {"strong", "good", "possible"}
    assert "extraction_confidence_level" in results[0]


def test_match_respects_preferences(client):
    manuscript_id = upload(client, docx_bytes()).json()["manuscript_id"]
    # The matcher counts hybrid journals as open-access capable.
    open_access = client.post(
        "/match", json={"manuscript_id": manuscript_id, "preferences": {"open_access_only": True}}
    ).json()
    assert {r["journal_id"] for r in open_access} == {"demo-ai-001", "demo-med-001"}

    cheap = client.post(
        "/match", json={"manuscript_id": manuscript_id, "preferences": {"max_apc": 1500}}
    ).json()
    assert {r["journal_id"] for r in cheap} == {"demo-ai-001", "demo-agri-001"}


def test_seed_does_not_overwrite_real_journals(tmp_path):
    s = Store(tmp_path / "real.db")
    s.save_journal(
        JournalRequirementSpec(
            journal_id="real", name="Real", publisher="P", source_url="https://r",
            hard_constraints=HardConstraint(), scope_description="s",
            extraction_confidence=0.9, access_model="hybrid",
        )
    )
    assert seed_if_empty(s) == 0
    assert [j.journal_id for j in s.list_journals()] == ["real"]
