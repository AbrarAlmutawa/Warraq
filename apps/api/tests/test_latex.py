"""LaTeX export: the current revision as a compilable article-class project."""

import io
import shutil
import subprocess
import zipfile
from pathlib import Path

import pytest
from docx import Document
from fastapi.testclient import TestClient

import main
from db import get_store
from db.seed import seed_if_empty
from db.store import Store
from services.analyzer.parser import parse_docx
from services.export.latex import escape, latex_zip, to_latex

DEMO = Path(__file__).parent / "fixtures" / "warraq_demo_manuscript.docx"


@pytest.fixture(scope="module")
def demo_tex():
    tex, figures, notes, xelatex = to_latex(DEMO.read_bytes(), parse_docx(DEMO))
    return tex, figures, notes, xelatex


def test_structure(demo_tex):
    tex, figures, notes, xelatex = demo_tex
    assert not xelatex and notes == []
    assert r"\title{Leakage-Safe Multimodal" in tex
    assert r"\author{Sara Example, Layla Sample, Noor Placeholder \\ Department of Demo Studies" in tex
    assert tex.index(r"\begin{abstract}") < tex.index(r"\textbf{Keywords:}") < tex.index(r"\end{abstract}")
    # LaTeX numbers sections itself, so the manual numbers are dropped.
    assert r"\section{Introduction}" in tex and "1. Introduction" not in tex
    assert r"\subsection{Dataset}" in tex
    assert r"\section*{Ethics Statement}" in tex


def test_tables_figures_and_references(demo_tex):
    tex, figures, _, _ = demo_tex
    assert tex.count(r"\begin{table}") == 2
    assert r"\caption{Class distribution of the wrist radiograph dataset.}" in tex
    assert list(figures) == ["figures/figure1.png"] and figures["figures/figure1.png"][:4] == b"\x89PNG"
    assert r"\includegraphics[width=0.8\linewidth]{figures/figure1.png}" in tex
    assert r"\caption{Accuracy and macro F1 score of the three models on the test set.}" in tex
    assert tex.count(r"\bibitem{") == 22
    assert r"\bibitem{ref1} S. Example and R. Sample" in tex  # "[1]" prefix removed
    assert r"\cite{ref1}" in tex and r"\cite{ref8}" in tex   # numeric citations linked


def test_escaping():
    assert escape("50% & $5_x {a} #1 ~ ^") == r"50\% \& \$5\_x \{a\} \#1 \textasciitilde{} \textasciicircum{}"
    assert escape("p ≤ 0.05 ± 2") == r"p $\leq$ 0.05 $\pm$ 2"
    assert escape("ok ≤", unicode_safe=False) == "ok ≤"


def test_author_year_citations_are_left_as_text():
    doc = Document()
    doc.add_paragraph("An APA Paper", style="Title")
    doc.add_paragraph("Introduction", style="Heading 1")
    doc.add_paragraph("As shown before (Example, 2023), results vary [sic].")
    doc.add_paragraph("References", style="Heading 1")
    doc.add_paragraph("Example, A. (2023). A study. Demo Journal, 1, 1-2.")
    doc.add_paragraph("Sample, B. (2021). Another study. Demo Journal, 2, 3-4.")
    buffer = io.BytesIO()
    doc.save(buffer)
    tex, _, _, _ = to_latex(buffer.getvalue(), parse_docx(io.BytesIO(buffer.getvalue())))
    assert r"\cite" not in tex and "(Example, 2023)" in tex
    assert tex.count(r"\bibitem{") == 2


def test_arabic_manuscripts_use_xelatex():
    doc = Document()
    doc.add_paragraph("تحليل المشاعر في مراجعات الأماكن السياحية", style="Title")
    doc.add_paragraph("المقدمة", style="Heading 1")
    doc.add_paragraph("تنمو السياحة بسرعة، وتكثر المراجعات على الإنترنت.")
    buffer = io.BytesIO()
    doc.save(buffer)
    tex, _, _, xelatex = to_latex(buffer.getvalue(), parse_docx(io.BytesIO(buffer.getvalue())))
    assert xelatex and r"\usepackage{polyglossia}" in tex and r"\setmainlanguage{arabic}" in tex
    assert tex.index(r"\usepackage[hidelinks]{hyperref}") < tex.index(r"\usepackage{polyglossia}")


@pytest.mark.skipif(shutil.which("pdflatex") is None, reason="pdflatex not installed")
def test_demo_compiles_with_pdflatex(tmp_path):
    with zipfile.ZipFile(io.BytesIO(latex_zip(DEMO.read_bytes(), parse_docx(DEMO), folder="demo"))) as archive:
        archive.extractall(tmp_path)
    folder = tmp_path / "demo"
    for _ in range(2):
        result = subprocess.run(["pdflatex", "-interaction=nonstopmode", "-halt-on-error", "main.tex"],
                                cwd=folder, capture_output=True, text=True, timeout=120)
    assert result.returncode == 0, result.stdout[-2000:]
    assert (folder / "main.pdf").exists()


# ------------------------------------------------------------------- API

@pytest.fixture
def client(tmp_path):
    store = Store(tmp_path / "latex.db")
    seed_if_empty(store)
    main.app.dependency_overrides[get_store] = lambda: store
    yield TestClient(main.app)
    main.app.dependency_overrides.clear()


def test_download_latex_follows_the_latest_revision(client):
    with DEMO.open("rb") as f:
        mid = client.post("/manuscripts/upload", files={"file": ("demo.docx", f)}).json()["manuscript_id"]
    client.patch(f"/manuscripts/{mid}/blocks/paragraph_0", json={"text": "A Shorter Title"})

    response = client.get(f"/manuscripts/{mid}/download", params={"format": "latex"})
    assert response.status_code == 200 and response.headers["content-type"] == "application/zip"
    assert "demo-warraq-v1-latex.zip" in response.headers["content-disposition"]
    with zipfile.ZipFile(io.BytesIO(response.content)) as archive:
        names = set(archive.namelist())
        assert {"demo-warraq-v1/main.tex", "demo-warraq-v1/README.txt", "demo-warraq-v1/figures/figure1.png"} <= names
        assert r"\title{A Shorter Title}" in archive.read("demo-warraq-v1/main.tex").decode()

    original = client.get(f"/manuscripts/{mid}/download", params={"format": "latex", "revision": 0})
    with zipfile.ZipFile(io.BytesIO(original.content)) as archive:
        assert r"\title{Leakage-Safe" in archive.read("demo/main.tex").decode()

    assert client.get(f"/manuscripts/{mid}/download", params={"format": "pdf"}).status_code == 422
