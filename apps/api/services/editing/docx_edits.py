"""
Edits applied to the researcher's real DOCX file (S4).

Why edit the Word file instead of plain text: the result keeps the original
fonts, styles, tables and figures, the parser re-reads it exactly as it read
the upload (so counts stay consistent), and the researcher can download a
file they can actually submit.

Block ids come from the parser ("paragraph_12" = doc.paragraphs[12]) and are
only valid for the revision they were parsed from. Every edit starts from the
latest revision and its own parse, so ids always line up.
"""

import io
import re
from copy import deepcopy

from docx import Document
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph

from services.analyzer.models import ManuscriptParsedData
from services.analyzer.parser import _heading_key, parse_docx

MAX_TEXT = 20_000


class EditError(ValueError):
    """The edit cannot be applied to this manuscript (shown to the researcher)."""


# ------------------------------------------------------------------ helpers

def load(docx: bytes):
    return Document(io.BytesIO(docx))


def save(doc) -> bytes:
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def reparse(docx: bytes) -> ManuscriptParsedData:
    return parse_docx(io.BytesIO(docx))


def _clean(text: str) -> str:
    text = (text or "").strip()
    if not text:
        raise EditError("The new text is empty.")
    if len(text) > MAX_TEXT:
        raise EditError("The new text is too long.")
    return text


def _index(block_id: str) -> int:
    match = re.fullmatch(r"paragraph_(\d+)", block_id or "")
    if not match:
        raise EditError(f"Unknown block '{block_id}'.")
    return int(match.group(1))


def _paragraph(doc, parsed: ManuscriptParsedData, block_id: str):
    if not any(b.id == block_id for b in parsed.blocks):
        raise EditError(f"Block '{block_id}' is not in the current version of the manuscript.")
    index = _index(block_id)
    paragraphs = doc.paragraphs
    if index >= len(paragraphs):
        raise EditError(f"Block '{block_id}' is not in the current version of the manuscript.")
    return paragraphs[index]


def set_text(paragraph, text: str) -> None:
    """Replace a paragraph's text, keeping the formatting of its first run."""
    runs = paragraph.runs
    if runs:
        runs[0].text = text
        for run in runs[1:]:
            run._element.getparent().remove(run._element)
    else:
        paragraph.add_run(text)


def delete(paragraph) -> None:
    paragraph._element.getparent().remove(paragraph._element)


def format_like(paragraph, template) -> None:
    """Give `paragraph` the paragraph and run formatting of `template` (None = document default)."""
    if paragraph._p.pPr is not None:
        paragraph._p.remove(paragraph._p.pPr)
    if template is not None and template._p.pPr is not None:
        paragraph._p.insert(0, deepcopy(template._p.pPr))
    template_run = template.runs[0] if template is not None and template.runs else None
    for run in paragraph.runs:
        if run._r.rPr is not None:
            run._r.remove(run._r.rPr)
        if template_run is not None and template_run._r.rPr is not None:
            run._r.insert(0, deepcopy(template_run._r.rPr))


def insert_after(paragraph, text: str, template=None):
    """New paragraph right after `paragraph`, formatted like `template` (default: `paragraph`)."""
    new = OxmlElement("w:p")
    paragraph._p.addnext(new)
    para = Paragraph(new, paragraph._parent)
    para.add_run(text)
    format_like(para, paragraph if template is None else template)
    return para


def insert_before(paragraph, text: str, template):
    para = paragraph.insert_paragraph_before(text)
    format_like(para, template)
    return para


def _style(doc, name: str):
    try:
        return doc.styles[name]
    except KeyError:
        return None


def _block_paragraphs(doc, parsed, block_ids):
    paragraphs = doc.paragraphs
    return [paragraphs[_index(b)] for b in block_ids if _index(b) < len(paragraphs)]


def _heading_block(parsed: ManuscriptParsedData, section: str):
    return next((b for b in parsed.blocks if b.type == "heading" and (b.section or "").lower() == section.lower()), None)


def _templates(doc, parsed):
    """Existing heading and body paragraphs to copy formatting from, so additions look native."""
    paragraphs = doc.paragraphs
    heading = next((b for b in parsed.blocks if b.type == "heading"), None)
    body = next((b for b in parsed.blocks if b.type == "paragraph" and b.section), None)
    return (paragraphs[heading.paragraph_index] if heading else None,
            paragraphs[body.paragraph_index] if body else None)


# ---------------------------------------------------------------- operations

def set_block_text(doc, parsed, block_id: str, text: str) -> str:
    paragraph = _paragraph(doc, parsed, block_id)
    set_text(paragraph, _clean(text))
    return f"Edited {block_id}"


def set_title(doc, parsed, text: str) -> str:
    title = next((b for b in parsed.blocks if b.type == "title"), None)
    if title is None:
        raise EditError("Warraq could not find the title in this manuscript.")
    set_text(doc.paragraphs[title.paragraph_index], _clean(text))
    return "Title updated"


def replace_abstract(doc, parsed, text: str) -> str:
    section = next((s for s in parsed.sections if s.name.lower() == "abstract"), None)
    if section is None or not section.block_ids:
        raise EditError("Warraq could not find the abstract in this manuscript.")
    parts = [p.strip() for p in re.split(r"\n\s*\n|\n", _clean(text)) if p.strip()]
    old = _block_paragraphs(doc, parsed, section.block_ids)
    set_text(old[0], parts[0])
    last = old[0]
    for part in parts[1:]:
        last = insert_after(last, part)
    for paragraph in old[1:]:
        delete(paragraph)
    return "Abstract updated"


def upsert_section(doc, parsed, heading: str, lines: list[str], *, position: str, bullets: bool = False) -> str:
    """
    Add a section (e.g. Highlights, Data Availability Statement), or replace the
    text of an existing section with the same heading so it is never duplicated.
    position: "before_abstract" or "before_references".
    """
    heading = _clean(heading)
    lines = [_clean(line) for line in lines if line and line.strip()]
    if not lines:
        raise EditError("The section has no text.")

    key = _heading_key(heading)
    existing = next((b for b in parsed.blocks if b.type == "heading" and _heading_key(b.text) == key), None)
    heading_template, body_template = _templates(doc, parsed)

    if existing is not None:
        section = next((s for s in parsed.sections if s.name == existing.section), None)
        for paragraph in _block_paragraphs(doc, parsed, section.block_ids if section else []):
            delete(paragraph)
        last = doc.paragraphs[existing.paragraph_index]
        for line in lines:
            last = insert_after(last, line, template=body_template)
            _bulletize(doc, last, bullets)
        return f"Updated {heading}"

    target_section = "Abstract" if position == "before_abstract" else "References"
    target = _heading_block(parsed, target_section)

    if target is not None:
        before = doc.paragraphs[target.paragraph_index]
        insert_before(before, heading, heading_template)
        for line in lines:
            _bulletize(doc, insert_before(before, line, body_template), bullets)
    else:
        if position == "before_abstract":
            title = next((b for b in parsed.blocks if b.type == "title"), None)
            anchor = doc.paragraphs[title.paragraph_index if title else 0]
        else:
            anchor = doc.paragraphs[-1]
        last = insert_after(anchor, heading, template=heading_template)
        for line in lines:
            last = insert_after(last, line, template=body_template)
            _bulletize(doc, last, bullets)
    return f"Added {heading}"


def _bulletize(doc, paragraph, bullets: bool) -> None:
    if not bullets:
        return
    bullet_style = _style(doc, "List Bullet")
    if bullet_style is not None:
        paragraph.style = bullet_style
    elif not paragraph.text.startswith("\u2022"):
        set_text(paragraph, f"\u2022 {paragraph.text}")


def replace_references(doc, parsed, references: list[str]) -> str:
    refs = [_clean(r) for r in references if r and r.strip()]
    if not refs:
        raise EditError("The reference list is empty.")
    blocks = [b for b in parsed.blocks if b.type == "reference"]
    if not blocks:
        raise EditError("Warraq could not find the reference list in this manuscript.")
    old = [doc.paragraphs[b.paragraph_index] for b in blocks]
    for paragraph, text in zip(old, refs):
        set_text(paragraph, text)
    last = old[min(len(old), len(refs)) - 1]
    for text in refs[len(old):]:
        last = insert_after(last, text)
    for paragraph in old[len(refs):]:
        delete(paragraph)
    return f"References replaced ({len(refs)})"


# ------------------------------------------------------------ direct editing

def _normalize(text: str) -> str:
    return re.sub(r"\s+", " ", text or "").strip()


def apply_text(doc, parsed, paragraphs: list[str]) -> str | None:
    """
    Make the document's paragraphs match what the researcher typed in the editor.

    `paragraphs` is the editor text split into lines (one paragraph per line, as the
    editor shows them). It is aligned with the current paragraphs, and only the
    differences are written to the Word file: changed paragraphs get their new text,
    new lines become new paragraphs (formatted like the paragraph before them, so a
    line typed after the last reference becomes a reference), and removed lines are
    deleted. Untouched paragraphs, tables and images are left exactly as they were.
    Returns None when nothing changed.
    """
    import difflib

    new = [_normalize(p) for p in paragraphs]
    new = [p for p in new if p]
    if any(len(p) > MAX_TEXT for p in new):
        raise EditError("One of the paragraphs is too long.")

    blocks, seen = [], set()
    for block in sorted(parsed.blocks, key=lambda b: b.paragraph_index):
        text = _normalize(block.text)
        if text and block.id not in seen:
            seen.add(block.id)
            blocks.append((block, text))
    if not new and blocks:
        raise EditError("The manuscript can't be left empty.")

    paragraphs_in_doc = doc.paragraphs
    old = [paragraphs_in_doc[b.paragraph_index] for b, _ in blocks]
    old_text = [t for _, t in blocks]

    matcher = difflib.SequenceMatcher(a=old_text, b=new, autojunk=False)
    changed = added = removed = 0
    for tag, i1, i2, j1, j2 in matcher.get_opcodes():
        if tag == "equal":
            continue
        if tag in ("replace", "delete"):
            pairs = min(i2 - i1, j2 - j1) if tag == "replace" else 0
            for k in range(pairs):
                set_text(old[i1 + k], new[j1 + k])
                changed += 1
            last = old[i1 + pairs - 1] if pairs else (old[i1 - 1] if i1 > 0 else None)
            for text in new[j1 + pairs:j2]:
                last = insert_after(last, text, template=old[i1]) if last is not None else insert_before(old[i1], text, old[i1])
                added += 1
            for paragraph in old[i1 + pairs:i2]:
                delete(paragraph)
                removed += 1
        elif tag == "insert":
            if i1 > 0:
                last = old[i1 - 1]
                for text in new[j1:j2]:
                    last = insert_after(last, text, template=old[i1 - 1])
                    added += 1
            elif old:
                for text in new[j1:j2]:
                    insert_before(old[0], text, old[0])
                    added += 1
            else:
                for text in new[j1:j2]:
                    doc.add_paragraph(text)
                    added += 1

    if not (changed or added or removed):
        return None
    parts = []
    if changed:
        parts.append(f"{changed} edited")
    if added:
        parts.append(f"{added} added")
    if removed:
        parts.append(f"{removed} removed")
    return "Text edited (" + ", ".join(parts) + ")"
