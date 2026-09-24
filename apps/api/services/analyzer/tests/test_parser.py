from pathlib import Path

from docx import Document

from services.analyzer.parser import parse_docx, count_words


def create_test_manuscript(path: Path):
    """
    Create a clean DOCX automatically for parser testing.
    """
    doc = Document()

    title = doc.add_paragraph("Test AI Manuscript")
    title.style = doc.styles["Title"]

    heading = doc.add_paragraph("Abstract")
    heading.style = doc.styles["Heading 1"]

    doc.add_paragraph(
        "Artificial intelligence supports academic writing."
    )

    doc.add_paragraph(
        "Keywords: AI; writing; research"
    )

    heading = doc.add_paragraph("1. Introduction")
    heading.style = doc.styles["Heading 1"]

    doc.add_paragraph(
        "This is the introduction section."
    )

    heading = doc.add_paragraph("2. Methodology")
    heading.style = doc.styles["Heading 1"]

    doc.add_paragraph(
        "This study used a simple experimental method."
    )

    heading = doc.add_paragraph("3. Results")
    heading.style = doc.styles["Heading 1"]

    doc.add_paragraph(
        "The experiment produced positive results."
    )

    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Group"
    table.cell(0, 1).text = "Score"
    table.cell(1, 0).text = "AI"
    table.cell(1, 1).text = "90"

    heading = doc.add_paragraph("References")
    heading.style = doc.styles["Heading 1"]

    doc.add_paragraph(
        "Smith, J. (2024). Example academic reference."
    )

    doc.add_paragraph(
        "Lee, A. (2025). Another academic reference."
    )

    doc.add_paragraph(
        "This is only a test note."
    )

    doc.save(path)


def create_messy_test_manuscript(path: Path):
    """
    Create a DOCX with inconsistent Word formatting.

    The section names should still be detected by the parser.
    """
    doc = Document()

    # Title deliberately uses Normal style
    doc.add_paragraph("Messy AI Manuscript")

    # Abstract deliberately uses Normal style
    doc.add_paragraph("Abstract")
    doc.add_paragraph(
        "Artificial intelligence supports academic writing."
    )

    # Non-standard keyword separator
    doc.add_paragraph(
        "Keywords - AI; writing; research"
    )

    # Plain-text section names
    doc.add_paragraph("Introduction")
    doc.add_paragraph(
        "This is the introduction section."
    )

    doc.add_paragraph("METHODS")
    doc.add_paragraph(
        "This study used a simple experimental method."
    )

    doc.add_paragraph("3 Results")
    doc.add_paragraph(
        "The experiment produced positive results."
    )

    # Mixed Word heading levels
    discussion = doc.add_paragraph("DISCUSSION")
    discussion.style = doc.styles["Heading 2"]

    doc.add_paragraph(
        "The findings were discussed here."
    )

    # Plain-text Conclusion
    doc.add_paragraph("Conclusion")
    doc.add_paragraph(
        "This is the final conclusion."
    )

    data_availability = doc.add_paragraph(
        "Data Availability Statement"
    )
    data_availability.style = doc.styles["Heading 3"]

    doc.add_paragraph(
        "Data are available upon reasonable request."
    )

    # References heading deliberately uses Normal style
    doc.add_paragraph("References")

    doc.add_paragraph(
        "Smith, J. (2024). Example academic reference."
    )

    doc.add_paragraph(
        "Lee, A. (2025). Another academic reference."
    )

    # Must not be counted as a reference
    doc.add_paragraph(
        "This is only a test note."
    )

    doc.save(path)


def test_count_words():
    assert count_words(
        "Artificial intelligence improves academic writing."
    ) == 5

    assert count_words(
        "يساعد الذكاء الاصطناعي الباحثين في الكتابة الأكاديمية."
    ) == 7


def test_parse_docx(tmp_path):
    test_file = tmp_path / "test_manuscript.docx"

    create_test_manuscript(test_file)

    result = parse_docx(test_file)

    assert result.title == "Test AI Manuscript"
    assert result.title_word_count == 3

    assert result.abstract == (
        "Artificial intelligence supports academic writing."
    )
    assert result.abstract_word_count == 5

    assert result.keywords == [
        "AI",
        "writing",
        "research",
    ]

    assert result.reference_count == 2
    assert len(result.references) == 2

    assert result.table_count == 1
    assert result.figure_count == 0

    section_names = [
        section.name
        for section in result.sections
    ]

    assert "Abstract" in section_names
    assert "Introduction" in section_names
    assert "Methodology" in section_names
    assert "Results" in section_names
    assert "References" in section_names

    references_section = next(
        section
        for section in result.sections
        if section.name == "References"
    )

    assert len(references_section.block_ids) == 2

    assert len(result.blocks) > 0

    abstract_blocks = [
        block
        for block in result.blocks
        if block.section == "Abstract"
        and block.type == "paragraph"
    ]

    assert len(abstract_blocks) == 1
    assert abstract_blocks[0].paragraph_index is not None


def test_parse_messy_docx(tmp_path):
    """
    Regression test:
    Warraq should parse a manuscript even when researchers
    do not use consistent Word heading styles.
    """
    test_file = tmp_path / "messy_manuscript.docx"

    create_messy_test_manuscript(test_file)

    result = parse_docx(test_file)

    assert result.title == "Messy AI Manuscript"

    assert result.abstract == (
        "Artificial intelligence supports academic writing."
    )
    assert result.abstract_word_count == 5

    assert result.keywords == [
        "AI",
        "writing",
        "research",
    ]

    section_names = [
        section.name
        for section in result.sections
    ]

    assert section_names == [
        "Abstract",
        "Introduction",
        "Methodology",
        "Results",
        "Discussion",
        "Conclusion",
        "Data Availability Statement",
        "References",
    ]

    assert result.reference_count == 2

    references_section = next(
        section
        for section in result.sections
        if section.name == "References"
    )

    assert len(references_section.block_ids) == 2

    # Confirm the parser normalized the messy section labels.
    methodology_heading = next(
        block
        for block in result.blocks
        if block.type == "heading"
        and block.section == "Methodology"
    )

    results_heading = next(
        block
        for block in result.blocks
        if block.type == "heading"
        and block.section == "Results"
    )

    assert methodology_heading.text == "METHODS"
    assert results_heading.text == "3 Results"
def test_parent_section_with_subsections(tmp_path):
    """
    A parent section should include the content of its
    numbered subsections instead of becoming an empty section.
    """
    test_file = tmp_path / "subsections.docx"

    doc = Document()

    title = doc.add_paragraph("Subsection Test Manuscript")
    title.style = doc.styles["Title"]

    methods = doc.add_paragraph("2. Methods")
    methods.style = doc.styles["Heading 1"]

    participants = doc.add_paragraph("2.1 Participants")
    participants.style = doc.styles["Heading 2"]

    doc.add_paragraph(
        "Twenty graduate students participated in the study."
    )

    procedure = doc.add_paragraph("2.2 Procedure")
    procedure.style = doc.styles["Heading 2"]

    doc.add_paragraph(
        "Participants completed two writing tasks during one session."
    )

    doc.save(test_file)

    result = parse_docx(test_file)

    section_names = [
        section.name
        for section in result.sections
    ]

    assert section_names == ["Methodology"]

    methodology = result.sections[0]

    assert methodology.word_count == 15

    assert "Twenty graduate students" in methodology.text
    assert "Participants completed two writing tasks" in methodology.text

    subsection_blocks = [
        block
        for block in result.blocks
        if block.type == "heading"
        and block.text in {
            "2.1 Participants",
            "2.2 Procedure",
        }
    ]

    assert len(subsection_blocks) == 2

    assert all(
        block.section == "Methodology"
        for block in subsection_blocks
    )
def test_table_reference_is_not_caption(tmp_path):
    """
    A normal sentence that starts with Table + number
    should not be classified as a table caption.
    """
    test_file = tmp_path / "caption_test.docx"

    doc = Document()

    title = doc.add_paragraph("Caption Test Manuscript")
    title.style = doc.styles["Title"]

    results = doc.add_paragraph("Results")
    results.style = doc.styles["Heading 1"]

    doc.add_paragraph(
        "Table 2 and Figure 1 report the main results."
    )

    doc.add_paragraph(
        "Table 2. Main results"
    )

    doc.add_paragraph(
        "Figure 1: Performance by group"
    )

    doc.save(test_file)

    result = parse_docx(test_file)

    normal_sentence = next(
        block
        for block in result.blocks
        if block.text
        == "Table 2 and Figure 1 report the main results."
    )

    table_caption = next(
        block
        for block in result.blocks
        if block.text == "Table 2. Main results"
    )

    figure_caption = next(
        block
        for block in result.blocks
        if block.text == "Figure 1: Performance by group"
    )

    assert normal_sentence.type == "paragraph"
    assert table_caption.type == "table_caption"
    assert figure_caption.type == "figure_caption"