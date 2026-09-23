from pathlib import Path
import re

from docx import Document

from models import ManuscriptParsedData, SectionInfo, Block


# Change this back to demo.docx after the messy-document test if you want.
SAMPLE_PATH = Path(__file__).parent / "samples" / "messy_demo.docx"


# -------------------------------------------------------------------
# Heading detection
# -------------------------------------------------------------------

HEADING_ALIASES = {
    "abstract": "Abstract",
    "الملخص": "Abstract",

    "introduction": "Introduction",
    "intro": "Introduction",
    "المقدمة": "Introduction",

    "method": "Methodology",
    "methods": "Methodology",
    "methodology": "Methodology",
    "materials and methods": "Methodology",
    "materials & methods": "Methodology",
    "المنهجية": "Methodology",
    "الطريقة": "Methodology",

    "result": "Results",
    "results": "Results",
    "النتائج": "Results",

    "discussion": "Discussion",
    "المناقشة": "Discussion",

    "conclusion": "Conclusion",
    "conclusions": "Conclusion",
    "الخاتمة": "Conclusion",

    "data availability statement": "Data Availability Statement",
    "data availability": "Data Availability Statement",
    "بيان إتاحة البيانات": "Data Availability Statement",

    "references": "References",
    "reference": "References",
    "bibliography": "References",
    "المراجع": "References",

    "acknowledgements": "Acknowledgements",
    "acknowledgments": "Acknowledgements",
    "شكر وتقدير": "Acknowledgements",
}


BODY_EXCLUDED_SECTIONS = {
    "abstract",
    "references",
    "bibliography",
    "data availability statement",
    "acknowledgements",
    "acknowledgments",
    "الملخص",
    "المراجع",
    "بيان إتاحة البيانات",
    "شكر وتقدير",
}


def count_words(text: str) -> int:
    """
    Count words in Arabic or English text.
    """
    words = re.findall(r"\b[\w'-]+\b", text, flags=re.UNICODE)
    return len(words)


def normalize_heading_name(text: str) -> str:
    """
    Remove common numbering from a heading.

    Examples:
    '1. Introduction' -> 'Introduction'
    '3 Results' -> 'Results'
    '3.1 Data Collection' -> 'Data Collection'
    """
    text = text.strip()

    cleaned = re.sub(
        r"^\s*\d+(?:\.\d+)*[\.\)]?\s*",
        "",
        text,
    )

    return cleaned.strip()


def _heading_key(text: str) -> str:
    """
    Normalize heading text for matching.
    """
    text = normalize_heading_name(text)
    text = text.strip().lower()

    # Remove harmless punctuation around a heading.
    text = re.sub(r"[:：\-–—]+$", "", text).strip()

    # Normalize repeated whitespace.
    text = re.sub(r"\s+", " ", text)

    return text


def detect_heading(paragraph) -> str | None:
    """
    Detect a section heading even when the researcher did not use
    Word's Heading styles.

    Priority:
    1. Known academic heading text, regardless of Word style.
    2. Any paragraph using a Word Heading style.
    """
    text = paragraph.text.strip()

    if not text:
        return None

    key = _heading_key(text)

    # Strong fallback for messy manuscripts:
    # recognize known academic section names by their text.
    if key in HEADING_ALIASES:
        return HEADING_ALIASES[key]

    style_name = (
        paragraph.style.name.lower()
        if paragraph.style
        else ""
    )

    # If Word explicitly says it is a heading, keep it even if it is
    # not one of our standard section names.
    if style_name.startswith("heading"):
        cleaned = normalize_heading_name(text).strip()
        return cleaned if cleaned else None

    return None


def is_keywords_line(text: str) -> bool:
    """
    Recognize common keyword-label formats.

    Examples:
    Keywords: AI; writing
    Keywords - AI; writing
    الكلمات المفتاحية: الذكاء الاصطناعي؛ البحث
    """
    return bool(
        re.match(
            r"^\s*(keywords?|الكلمات\s+المفتاحية)\s*[:：\-–—]",
            text,
            flags=re.IGNORECASE,
        )
    )


# -------------------------------------------------------------------
# Core extraction
# -------------------------------------------------------------------

def extract_title(doc: Document) -> str:
    """
    Extract the manuscript title.

    First use Word's Title style.
    Fallback to the first non-empty paragraph.
    """
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()

        if (
            text
            and paragraph.style
            and paragraph.style.name.lower() == "title"
        ):
            return text

    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()

        if text:
            return text

    return ""


def extract_abstract(doc: Document) -> str:
    """
    Extract the abstract even when 'Abstract' is not using Heading 1.
    Stop when the next recognized section heading begins.
    """
    collecting = False
    abstract_parts = []

    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()

        if not text:
            continue

        detected_heading = detect_heading(paragraph)

        if detected_heading == "Abstract":
            collecting = True
            continue

        if collecting and detected_heading is not None:
            break

        if collecting:
            if is_keywords_line(text):
                continue

            abstract_parts.append(text)

    return "\n".join(abstract_parts)


def extract_keywords(doc: Document) -> list[str]:
    """
    Extract manuscript keywords from common label formats.
    """
    pattern = re.compile(
        r"^\s*(?:keywords?|الكلمات\s+المفتاحية)\s*"
        r"[:：\-–—]\s*(.+)$",
        flags=re.IGNORECASE,
    )

    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()

        if not text:
            continue

        match = pattern.match(text)

        if match:
            keyword_text = match.group(1).strip()

            keywords = re.split(
                r"[;؛,،]",
                keyword_text,
            )

            return [
                keyword.strip()
                for keyword in keywords
                if keyword.strip()
            ]

    return []


def extract_sections(doc: Document) -> dict[str, str]:
    """
    Extract sections using both Word styles and text-based fallback
    heading detection.
    """
    sections = {}
    current_section = None
    current_parts = []

    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()

        if not text:
            continue

        detected_heading = detect_heading(paragraph)

        if detected_heading is not None:
            if current_section is not None:
                sections[current_section] = "\n".join(
                    current_parts
                ).strip()

            current_section = detected_heading
            current_parts = []
            continue

        if current_section is not None:
            if (
                current_section == "Abstract"
                and is_keywords_line(text)
            ):
                continue

            current_parts.append(text)

    if current_section is not None:
        sections[current_section] = "\n".join(
            current_parts
        ).strip()

    return sections


# -------------------------------------------------------------------
# References
# -------------------------------------------------------------------

def looks_like_reference(text: str) -> bool:
    """
    Check whether a paragraph looks like a bibliography reference.
    Supports common author-year and numbered formats.
    """
    text = text.strip()

    if not text:
        return False

    has_year = re.search(
        r"\((?:19|20)\d{2}[a-z]?\)",
        text,
        flags=re.IGNORECASE,
    )

    numbered_brackets = re.match(
        r"^\s*\[\d+\]\s+",
        text,
    )

    numbered_list = re.match(
        r"^\s*\d+\.\s+",
        text,
    )

    return bool(
        has_year
        or numbered_brackets
        or numbered_list
    )


def extract_references(doc: Document) -> list[str]:
    """
    Extract individual references from the References section.
    """
    references = []
    collecting = False

    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()

        if not text:
            continue

        detected_heading = detect_heading(paragraph)

        if detected_heading == "References":
            collecting = True
            continue

        if collecting and detected_heading is not None:
            break

        if collecting and looks_like_reference(text):
            references.append(text)

    return references


# -------------------------------------------------------------------
# Counts
# -------------------------------------------------------------------

def count_tables(doc: Document) -> int:
    """
    Count actual Word tables.
    """
    return len(doc.tables)


def count_figures(doc: Document) -> int:
    """
    Count inline Word images/shapes.
    Good enough for the Warraq MVP.
    """
    return len(doc.inline_shapes)


def calculate_main_text_word_count(
    sections: dict[str, str]
) -> int:
    """
    Count the main research body while excluding metadata sections.
    """
    total = 0

    for section_name, section_text in sections.items():
        normalized_name = section_name.strip().lower()

        if normalized_name in BODY_EXCLUDED_SECTIONS:
            continue

        total += count_words(section_text)

    return total


# -------------------------------------------------------------------
# Traceable blocks
# -------------------------------------------------------------------

def build_blocks(doc: Document) -> list[Block]:
    """
    Convert manuscript paragraphs into traceable blocks.

    Every block preserves the original Word paragraph index.
    """
    blocks = []
    current_section = None
    title = extract_title(doc)
    title_assigned = False

    for index, paragraph in enumerate(doc.paragraphs):
        text = paragraph.text.strip()

        if not text:
            continue

        detected_heading = detect_heading(paragraph)

        if detected_heading is not None:
            current_section = detected_heading

            blocks.append(
                Block(
                    id=f"paragraph_{index}",
                    type="heading",
                    text=text,
                    section=current_section,
                    word_count=count_words(text),
                    paragraph_index=index,
                )
            )
            continue

        # Title can still be detected when its Word style is Normal.
        if not title_assigned and text == title:
            block_type = "title"
            title_assigned = True

        elif is_keywords_line(text):
            block_type = "keywords"

        elif re.match(
            r"^(table|جدول)\s*\d+",
            text,
            flags=re.IGNORECASE,
        ):
            block_type = "table_caption"

        elif re.match(
            r"^(figure|fig\.?|شكل)\s*\d+",
            text,
            flags=re.IGNORECASE,
        ):
            block_type = "figure_caption"

        elif (
            current_section == "References"
            and looks_like_reference(text)
        ):
            block_type = "reference"

        else:
            block_type = "paragraph"

        blocks.append(
            Block(
                id=f"paragraph_{index}",
                type=block_type,
                text=text,
                section=current_section,
                word_count=count_words(text),
                paragraph_index=index,
            )
        )

    return blocks


# -------------------------------------------------------------------
# Main parser contract
# -------------------------------------------------------------------

def parse_docx(path) -> ManuscriptParsedData:
    """
    Parse a DOCX manuscript and return structured Warraq data.
    """
    doc = Document(path)

    full_text = "\n".join(
        paragraph.text.strip()
        for paragraph in doc.paragraphs
        if paragraph.text.strip()
    )

    title = extract_title(doc)
    abstract = extract_abstract(doc)
    keywords = extract_keywords(doc)

    raw_sections = extract_sections(doc)

    # Clean References so non-reference trailing notes are excluded.
    references = extract_references(doc)

    if "References" in raw_sections:
        raw_sections["References"] = "\n".join(references)

    blocks = build_blocks(doc)

    section_block_ids = {}

    for block in blocks:
        if not block.section:
            continue

        if block.type in {"heading", "keywords"}:
            continue

        if (
            block.section == "References"
            and block.type != "reference"
        ):
            continue

        section_block_ids.setdefault(
            block.section,
            [],
        ).append(block.id)

    sections = [
        SectionInfo(
            name=section_name,
            text=section_text,
            word_count=count_words(section_text),
            block_ids=section_block_ids.get(
                section_name,
                [],
            ),
        )
        for section_name, section_text in raw_sections.items()
    ]

    return ManuscriptParsedData(
        title=title,
        title_word_count=count_words(title),

        abstract=abstract,
        abstract_word_count=count_words(abstract),

        keywords=keywords,

        full_document_word_count=count_words(full_text),
        main_text_word_count=calculate_main_text_word_count(
            raw_sections
        ),

        reference_count=len(references),
        references=references,

        table_count=count_tables(doc),
        figure_count=count_figures(doc),

        sections=sections,
        blocks=blocks,

        full_text=full_text,
    )


# -------------------------------------------------------------------
# Debugging
# -------------------------------------------------------------------

def inspect_docx(path=SAMPLE_PATH):
    doc = Document(path)

    print("\nPARAGRAPHS")
    print("=" * 80)

    for index, paragraph in enumerate(doc.paragraphs):
        text = paragraph.text.strip()

        if text:
            print(
                f"{index:03d} | "
                f"{paragraph.style.name:<20} | "
                f"{text[:120]}"
            )

    print("\nDOCUMENT INFO")
    print("=" * 80)
    print("Tables:", len(doc.tables))
    print("Images:", len(doc.inline_shapes))


if __name__ == "__main__":
    inspect_docx()

    parsed = parse_docx(SAMPLE_PATH)

    print("\nPARSER SUMMARY")
    print("=" * 80)
    print("Title:", parsed.title)
    print("Title words:", parsed.title_word_count)
    print("Abstract words:", parsed.abstract_word_count)
    print("Keywords:", parsed.keywords)
    print("Main text words:", parsed.main_text_word_count)
    print("References:", parsed.reference_count)
    print("Tables:", parsed.table_count)
    print("Figures:", parsed.figure_count)
    print(
        "Sections:",
        [section.name for section in parsed.sections],
    )

    print("\nFINAL PARSED DATA")
    print("=" * 80)
    print(parsed.model_dump_json(indent=2))
