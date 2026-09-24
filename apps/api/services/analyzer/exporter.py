from pathlib import Path
import re

from docx import Document
from docx.enum.text import WD_COLOR_INDEX
try:
    from .models import ValidationIssue
except ImportError:
    from models import ValidationIssue

SAMPLE_PATH = Path(__file__).parent / "samples" / "messy_demo.docx"

OUTPUT_PATH = (
    Path(__file__).parent
    / "samples"
    / "warraq_reviewed.docx"
)


def get_paragraph_index(block_id: str) -> int:
    """
    Convert a Warraq block ID into its Word paragraph index.

    Example:
    paragraph_4 -> 4
    """

    match = re.match(
        r"^paragraph_(\d+)$",
        block_id
    )

    if not match:
        raise ValueError(
            f"Invalid block ID: {block_id}"
        )

    return int(match.group(1))


def highlight_paragraph(paragraph):
    """
    Highlight all text runs inside one Word paragraph.
    """

    for run in paragraph.runs:
        run.font.highlight_color = (
            WD_COLOR_INDEX.YELLOW
        )


def create_highlighted_docx(
    input_path,
    output_path,
    issues: list[ValidationIssue | dict],
):
    """
    Create an editable copy of the original DOCX and
    highlight paragraphs linked to failed/warning
    journal validation issues.
    """

    doc = Document(input_path)

    highlighted_blocks = set()

    for raw_issue in issues:

        # Accept either a ValidationIssue object
        # or a normal dictionary from an API response
        issue = (
            raw_issue
            if isinstance(raw_issue, ValidationIssue)
            else ValidationIssue.model_validate(
                raw_issue
            )
        )

        # Passed rules do not need highlighting
        if issue.status == "pass":
            continue

        for block_id in issue.block_ids:

            if block_id in highlighted_blocks:
                continue

            paragraph_index = (
                get_paragraph_index(block_id)
            )

            if paragraph_index >= len(doc.paragraphs):
                continue

            paragraph = doc.paragraphs[
                paragraph_index
            ]

            highlight_paragraph(paragraph)

            highlighted_blocks.add(block_id)

    doc.save(output_path)

    return Path(output_path)


if __name__ == "__main__":

    demo_issues = [
    {
        "rule": "abstract_max_words",
        "status": "fail",
        "actual": 280,
        "limit": 250,
        "message": (
            "Abstract exceeds the journal "
            "limit by 30 words."
        ),
        "block_ids": [
            "paragraph_4"
        ],
        "source_url": (
            "https://example-journal.com/"
            "author-guidelines"
        ),
    },
    {
        "rule": "maximum_tables",
        "status": "pass",
        "actual": 2,
        "limit": 4,
        "message": (
            "Table count is within "
            "the journal limit."
        ),
        "block_ids": [
            "paragraph_14",
            "paragraph_20",
        ],
    },
]

    result_path = create_highlighted_docx(
        SAMPLE_PATH,
        OUTPUT_PATH,
        demo_issues,
    )

    print(
        "Reviewed manuscript created:"
    )

    print(result_path)